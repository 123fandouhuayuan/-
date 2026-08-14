"""
机床Excel导入模块

功能：解析机床评分Excel模板，提取机床信息和检测数据，保存到MySQL和达梦数据库

支持机床类型：加工中心、镗铣床
支持模板格式：新模板（规范工作表名称）、旧模板（Sheet1/Sheet2）

主要函数：
    - process_excel_import: 主入口函数
    - _process_excel_import_internal: 内部实现
    - update_machine_parameters_by_type: 根据机床类型更新参数
"""
import os
import re
import tempfile
from datetime import datetime
from openpyxl import load_workbook
from django.db import transaction
from django.utils import timezone
import logging

logger = logging.getLogger(__name__)


# =============================================================================
# 全局常量：字段映射定义
# =============================================================================

GEOMETRY_FIELD_MAPPING = {
    'X轴直线度（X-Y垂直平面内）': 'x_axis_linearity_xy',
    'X轴直线度（Z-X水平平面内）': 'x_axis_linearity_zx',
    'Z轴直线度（Y-Z垂直平面内）': 'z_axis_linearity_yz',
    'Z轴直线度（Z-X水平平面内）': 'z_axis_linearity_zx',
    'Y轴直线度（X-Y垂直平面内）': 'y_axis_linearity_xy',
    'Y轴直线度（Y-Z垂直平面内）': 'y_axis_linearity_yz',
    'X轴角度偏差（俯仰）': 'x_axis_angle_pitch',
    'X轴角度偏差（偏摆）': 'x_axis_angle_yaw',
    'X轴角度偏差（倾斜）': 'x_axis_angle_roll',
    'Z轴角度偏差（俯仰）': 'z_axis_angle_pitch',
    'Z轴角度偏差（偏摆）': 'z_axis_angle_yaw',
    'Z轴角度偏差（倾斜）': 'z_axis_angle_roll',
    'Y轴角度偏差（Y-Z垂直平面内）': 'y_axis_angle_yz',
    'Y轴角度偏差（X-Y垂直平面内）': 'y_axis_angle_xy',
    'Y轴角度偏差（ZX水平面内）': 'y_axis_angle_zx',
    'Y轴与X轴垂直度': 'y_axis_perpendicularity_x',
    'Y轴与Z轴垂直度': 'y_axis_perpendicularity_z',
    'Z轴与X轴垂直度': 'z_axis_perpendicularity_x',
    '主轴轴向窜动': 'spindle_axial_runout',
    '主轴锥孔径向跳动（靠近端部）': 'spindle_radial_runout_near',
    '主轴锥孔径向跳动（距端部300mm）': 'spindle_radial_runout_300mm',
    '主轴轴线与Z轴平行度（Y-Z平面）': 'spindle_parallelism_yz',
    '主轴轴线与Z轴平行度（Z-X平面）': 'spindle_parallelism_zx',
    '主轴轴线与X轴垂直度': 'spindle_perpendicularity_x',
    '主轴轴线与Y轴垂直度': 'spindle_perpendicularity_y',
    '工作台面平面度': 'worktable_flatness',
    '工作台面与X轴平行度': 'worktable_parallelism_x',
    '工作台面与Z轴平行度': 'worktable_parallelism_z'
}

BM_GEOMETRY_FIELD_MAPPING = {
    '立柱直线度（Y-Z垂直平面内）': 'bm_column_linearity_yz',
    '立柱直线度（Z-X水平平面内）': 'bm_column_linearity_zx',
    '立柱Z轴直线度（Y-Z垂直平面内）': 'bm_column_linearity_yz',
    '立柱Z轴直线度（Z-X水平平面内）': 'bm_column_linearity_zx',
    'X轴直线度（X-Y垂直平面内）': 'bm_worktable_x_linearity_xy',
    'X轴直线度（Z-X水平平面内）': 'bm_worktable_x_linearity_zx',
    'Y轴直线度（Y-Z平面内）': 'bm_spindle_box_y_linearity_yz',
    'Y轴直线度（X-Y平面内）': 'bm_spindle_box_y_linearity_xy',
    '立柱角度偏差（YZ俯仰）': 'bm_column_angle_yz',
    '立柱角度偏差（XY倾斜）': 'bm_column_angle_xy',
    '立柱角度偏差（ZX偏摆）': 'bm_column_angle_zx',
    'X轴角度偏差（YZ俯仰）': 'bm_worktable_x_angle_yz',
    'X轴角度偏差（XY俯仰）': 'bm_worktable_x_angle_xy',
    'X轴角度偏差（YZ倾斜）': 'bm_worktable_x_angle_yz',
    'X轴角度偏差（ZX偏摆）': 'bm_worktable_x_angle_zx',
    'X轴角度偏差（XY倾斜）': 'bm_worktable_x_angle_xy',
    'Y轴角度偏差（YZ俯仰）': 'bm_spindle_box_y_angle_yz',
    'Y轴角度偏差（ZX偏摆）': 'bm_spindle_box_y_angle_zx',
    'Y轴角度偏差（ZX倾斜）': 'bm_spindle_box_y_angle_zx',
    'X轴与立柱垂直度': 'bm_worktable_x_column_perpendicularity',
    'X轴与立柱轴垂直度': 'bm_worktable_x_column_perpendicularity',
    'Y轴与X轴垂直度': 'bm_spindle_box_y_worktable_x_perpendicularity',
    'Y轴与立柱垂直度': 'bm_spindle_box_y_column_perpendicularity',
    'Y轴与立柱轴垂直度': 'bm_spindle_box_y_column_perpendicularity',
    '工作台面平面度': 'bm_worktable_flatness',
    '工作台面的平面度': 'bm_worktable_flatness',
    '立柱移动的平行度': 'bm_column_parallelism',
    '工作台面对立柱移动的平行度': 'bm_column_parallelism',
    'X轴移动的平行度': 'bm_worktable_x_parallelism',
    '工作台面对X轴的平行度': 'bm_worktable_x_parallelism',
    '镗轴锥孔径向跳动（靠近端部）': 'bm_boring_axis_taper_radial_near',
    '镗轴锥孔的径向跳动靠近轴端部处': 'bm_boring_axis_taper_radial_near',
    '镗轴锥孔径向跳动（距端部300mm）': 'bm_boring_axis_taper_radial_300mm',
    '镗轴锥孔的径向跳动距镗轴端部300mm处': 'bm_boring_axis_taper_radial_300mm',
    '镗轴径向跳动（缩回）': 'bm_boring_axis_radial_retracted',
    '镗轴的径向跳动镗轴缩回时': 'bm_boring_axis_radial_retracted',
    '镗轴径向跳动（伸出300mm）': 'bm_boring_axis_radial_extended',
    '镗轴的径向跳动镗轴伸出300mm时': 'bm_boring_axis_radial_extended',
    '镗轴轴向窜动（缩回）': 'bm_boring_axis_axial_retracted',
    '镗轴周期性的轴向窜动': 'bm_boring_axis_axial_retracted',
    '铣轴径向跳动': 'bm_milling_axis_radial',
    '铣轴端部的径向跳动': 'bm_milling_axis_radial',
    '铣轴轴向窜动': 'bm_milling_axis_axial',
    '铣轴端部的周期性轴向窜动': 'bm_milling_axis_axial',
}


# =============================================================================
# 工具函数
# =============================================================================

def parse_year_month(year_month_str):
    """
    解析年月日期格式
    
    支持格式：YYYY-MM、YYYY-MM-DD、datetime对象、Excel日期数字
    返回：date对象，失败返回None
    """
    if not year_month_str:
        return None
    try:
        # 处理datetime对象（openpyxl读取Excel日期单元格时返回）
        if isinstance(year_month_str, datetime):
            return year_month_str.date() if isinstance(year_month_str, datetime) else year_month_str
        # 处理date对象
        import datetime as dt
        if isinstance(year_month_str, dt.date):
            return year_month_str
        # 处理数字类型（Excel日期序列号）
        if isinstance(year_month_str, (int, float)):
            try:
                from openpyxl.utils import get_column_letter
                # Excel日期序列号转换（以1900-01-01为基准）
                from datetime import timedelta
                base_date = dt.date(1899, 12, 30)
                return base_date + timedelta(days=int(year_month_str))
            except Exception:
                return None
        # 处理字符串格式
        year_month_str = str(year_month_str).strip()
        if len(year_month_str) == 7 and '-' in year_month_str:
            return datetime.strptime(f"{year_month_str}-01", '%Y-%m-%d').date()
        elif len(year_month_str) == 10:
            return datetime.strptime(year_month_str, '%Y-%m-%d').date()
        # 处理 YYYY/MM 或 YYYY/MM/DD 格式
        elif '/' in year_month_str:
            year_month_str_normalized = year_month_str.replace('/', '-')
            if len(year_month_str_normalized) == 7:
                return datetime.strptime(f"{year_month_str_normalized}-01", '%Y-%m-%d').date()
            elif len(year_month_str_normalized) == 10:
                return datetime.strptime(year_month_str_normalized, '%Y-%m-%d').date()
        return None
    except Exception:
        return None


def extract_number(value):
    """
    从字符串中提取数值
    
    支持格式：纯数字、带小数、带单位（如"320 mm"）、带符号
    返回：int或float，失败返回None
    """
    if value is None or value == '':
        return None
    value_str = str(value).strip()
    match = re.search(r'[-+]?[0-9]*\.?[0-9]+', value_str)
    if match:
        try:
            num_str = match.group()
            if '.' in num_str:
                return float(num_str)
            else:
                return int(num_str)
        except (ValueError, TypeError):
            return None
    return None


def extract_range(value):
    """
    从字符串中提取范围值（上下限）
    
    支持格式："370~390"、"380±10"、"370/390"、单个数值
    返回：(min_value, max_value) 元组，失败返回 (None, None)
    """
    if value is None or value == '':
        return (None, None)
    
    value_str = str(value).strip()
    
    match = re.search(r'([-+]?[0-9]*\.?[0-9]+)\s*[~\-–—/]\s*([-+]?[0-9]*\.?[0-9]+)', value_str)
    if match:
        try:
            min_val = float(match.group(1)) if '.' in match.group(1) else int(match.group(1))
            max_val = float(match.group(2)) if '.' in match.group(2) else int(match.group(2))
            return (min(min_val, max_val), max(min_val, max_val))
        except (ValueError, TypeError):
            pass
    
    match = re.search(r'([-+]?[0-9]*\.?[0-9]+)\s*±\s*([-+]?[0-9]*\.?[0-9]+)', value_str)
    if match:
        try:
            center = float(match.group(1)) if '.' in match.group(1) else int(match.group(1))
            delta = float(match.group(2)) if '.' in match.group(2) else int(match.group(2))
            return (center - delta, center + delta)
        except (ValueError, TypeError):
            pass
    
    numbers = re.findall(r'[-+]?[0-9]*\.?[0-9]+', value_str)
    if len(numbers) >= 2:
        try:
            val1 = float(numbers[0]) if '.' in numbers[0] else int(numbers[0])
            val2 = float(numbers[1]) if '.' in numbers[1] else int(numbers[1])
            return (min(val1, val2), max(val1, val2))
        except (ValueError, TypeError):
            pass
    
    single_value = extract_number(value_str)
    if single_value is not None:
        return (single_value, single_value)
    
    return (None, None)


def update_machine_parameters_by_type(DamengService, machine_info, update_data, request):
    """
    根据机床类型选择对应的参数更新方法
    
    参数：
        DamengService: 达梦数据库服务类
        machine_info: 机床信息字典，包含id和machine_type
        update_data: 要更新的数据字典
        request: HTTP请求对象
    
    返回：更新结果字典
    """
    if machine_info.get('machine_type') == 'boring_milling_machine':
        return DamengService.update_boring_milling_parameters(
            machine_id=machine_info['id'],
            update_data=update_data,
            user=request.user.username,
            ip_address=request.META.get('REMOTE_ADDR', '127.0.0.1')
        )
    else:
        return DamengService.update_machine_parameters(
            machine_id=machine_info['id'],
            update_data=update_data,
            user=request.user.username,
            ip_address=request.META.get('REMOTE_ADDR', '127.0.0.1')
        )

def process_excel_import(data_file, request, Machine):
    """
    处理Excel文件导入的主入口函数
    
    参数：
        data_file: 上传的Excel文件对象
        request: HTTP请求对象
        Machine: 机床模型类
    
    返回：包含success、message、imported_count等信息的字典
    
    说明：每台机床使用独立savepoint，单台失败不影响其他机床
    """
    from apps.data_management.services.dameng_service import DamengDatabaseService as DamengService
    
    # 不使用外层atomic，让每台机床使用独立savepoint，确保单台机床失败不影响其他机床
    return _process_excel_import_internal(data_file, request, Machine, DamengService)


def _process_excel_import_internal(data_file, request, Machine, DamengService):
    """
    处理Excel文件导入的内部实现
    
    参数：
        data_file: 上传的Excel文件对象
        request: HTTP请求对象
        Machine: 机床模型类
        DamengService: 达梦数据库服务类
    
    返回：包含success、message、imported_count等信息的字典
    
    处理流程：
        1. 保存临时文件并加载Excel
        2. 检测模板类型（新/旧）
        3. 解析机床信息并创建记录
        4. 解析各模块数据并保存到达梦数据库
        5. 清理临时文件
    """
    error_messages = []
    
    file_name = data_file.name
    file_extension = os.path.splitext(file_name)[1].lower()
    
    with tempfile.NamedTemporaryFile(suffix=file_extension, delete=False) as temp_file:
        for chunk in data_file.chunks():
            temp_file.write(chunk)
        temp_file_path = temp_file.name
    
    try:
        # 打开Excel文件
        input_workbook = load_workbook(temp_file_path, data_only=True)
        
        # 检查工作表数量
        if len(input_workbook.sheetnames) < 2:
            return {'success': False, 'message': 'Excel文件至少需要包含2个工作表'}
        
        # 使用工作表索引获取第一个工作表
        machine_sheet = input_workbook.worksheets[0]  # 第一个工作表
        
        logger.info(f'第一个工作表名称: {machine_sheet.title}')
        
        # 检测模板类型
        is_new_template = machine_sheet.title == '机床信息'
        logger.info(f'模板类型: {"新模板" if is_new_template else "旧模板"}')
        
        if is_new_template:
            # 新模板处理逻辑
            # 读取第一个工作表表头（第1行）
            machine_header_row_1 = list(machine_sheet.iter_rows(min_row=1, max_row=1, values_only=True))[0]
            logger.info(f'第一个工作表表头1: {machine_header_row_1}')
            
            # 读取第二个表头（第2行）
            machine_header_row_2 = list(machine_sheet.iter_rows(min_row=2, max_row=2, values_only=True))[0]
            logger.info(f'第一个工作表表头2: {machine_header_row_2}')
            
            # 验证第一个工作表表头是否包含必需字段
            required_machine_fields = ['机床名称', '报告编号', '机床类型']
            missing_machine_fields = [field for field in required_machine_fields if field not in machine_header_row_1]
            if missing_machine_fields:
                return {'success': False, 'message': f'第一个工作表缺少必需字段: {", ".join(missing_machine_fields)}'}
            
            # 读取第一个工作表数据（从第3行开始，跳过两行表头）
            machine_rows = list(machine_sheet.iter_rows(min_row=3, values_only=True))
            logger.info(f'第一个工作表包含 {len(machine_rows)} 行数据')
            
            # 验证第一个工作表数据
            valid_machine_rows = []
            for row_idx, row in enumerate(machine_rows, start=3):
                # 跳过完全空行
                if not any(cell for cell in row):
                    continue
                
                # 验证必填字段（新模板列位置）
                # 第0列: 机床类型, 第1列: 机床名称, 第3列: 报告编号
                machine_type = row[0] if len(row) > 0 else ''
                machine_name = row[1] if len(row) > 1 else ''
                report_number = row[3] if len(row) > 3 else ''
                
                if not machine_name or not report_number:
                    error_msg = f'机床信息第{row_idx}行：机床名称或报告编号为空，跳过该行'
                    logger.warning(error_msg)
                    error_messages.append(error_msg)
                    continue
                
                valid_machine_rows.append(row)
            
            machine_rows = valid_machine_rows
            logger.info(f'机床信息有效数据行数: {len(machine_rows)}')
            
            # 检测是否为多机床模板
            if len(machine_rows) > 1:
                logger.info(f'检测到多机床模板，共 {len(machine_rows)} 台机床待导入')
            else:
                logger.info('检测到单机床模板')
        else:
            # 旧模板处理逻辑
            # 读取第一个工作表表头
            machine_header_row = list(machine_sheet.iter_rows(min_row=1, max_row=1, values_only=True))[0]
            logger.info(f'第一个工作表表头: {machine_header_row}')
            
            # 验证第一个工作表表头是否包含必需字段
            required_machine_fields = ['机床名称', '报告编号', '机床类型']
            missing_machine_fields = [field for field in required_machine_fields if field not in machine_header_row]
            if missing_machine_fields:
                return {'success': False, 'message': f'第一个工作表缺少必需字段: {", ".join(missing_machine_fields)}'}
            
            # 读取第一个工作表数据（从第2行开始）
            machine_rows = list(machine_sheet.iter_rows(min_row=2, values_only=True))
            logger.info(f'第一个工作表包含 {len(machine_rows)} 行数据')
            
            # 验证第一个工作表数据
            valid_machine_rows = []
            for row_idx, row in enumerate(machine_rows, start=2):
                # 跳过完全空行
                if not any(cell for cell in row):
                    continue
                
                # 验证必填字段
                machine_name = row[1] if len(row) > 1 else ''
                report_number = row[2] if len(row) > 2 else ''
                
                if not machine_name or not report_number:
                    error_msg = f'Sheet1第{row_idx}行：机床名称或报告编号为空，跳过该行'
                    logger.warning(error_msg)
                    error_messages.append(error_msg)
                    continue
                
                valid_machine_rows.append(row)
            
            machine_rows = valid_machine_rows
            logger.info(f'Sheet1有效数据行数: {len(machine_rows)}')
        
        if len(machine_rows) == 0:
            return {'success': False, 'message': '机床信息中没有有效的机床数据，请检查机床名称和报告编号是否填写完整'}
        
        # 根据模板类型处理检测项目数据
        detection_rows = []
        field_mapping = {}
        
        if not is_new_template:
            # 旧模板处理逻辑：读取第二个工作表
            detection_sheet = input_workbook.worksheets[1]  # 第二个工作表
            logger.info(f'第二个工作表名称: {detection_sheet.title}')
            
            # 读取第二个工作表表头
            detection_header_row = list(detection_sheet.iter_rows(min_row=1, max_row=1, values_only=True))[0]
            logger.info(f'第二个工作表表头: {detection_header_row}')
            
            # 验证第二个工作表表头是否包含"实测值"列
            if not any('实测值' in str(cell) for cell in detection_header_row):
                return {'success': False, 'message': '第二个工作表缺少"实测值"列'}
            
            # 读取第二个工作表数据（从第2行到第29行）
            detection_rows = list(detection_sheet.iter_rows(min_row=2, max_row=29, values_only=True))
            logger.info(f'第二个工作表包含 {len(detection_rows)} 行数据')
            
            # 验证第二个工作表数据
            valid_detection_rows = []
            for row_idx, row in enumerate(detection_rows, start=2):
                # 跳过完全空行
                if not any(cell for cell in row):
                    continue
                
                # 验证检测项目描述（第3列，索引为2）
                item_description = row[2] if len(row) > 2 else None
                if not item_description or not str(item_description).strip():
                    error_msg = f'Sheet2第{row_idx}行：检测项目描述为空，跳过该行'
                    logger.warning(error_msg)
                    continue
                
                valid_detection_rows.append(row)
            
            detection_rows = valid_detection_rows
            logger.info(f'Sheet2有效数据行数: {len(detection_rows)}')
            
            # 调试：打印前3行数据
            if len(detection_rows) > 0:
                logger.info(f'Sheet2前3行数据示例:')
                for i, row in enumerate(detection_rows[:3]):
                    logger.info(f'  第{i+2}行: {row}')
            
            if len(detection_rows) == 0:
                return {'success': False, 'message': 'Sheet2中没有有效的检测项目数据，请检查检测项目描述是否填写完整'}
            
            logger.info(f'使用全局字段映射：加工中心 {len(GEOMETRY_FIELD_MAPPING)} 个，镗铣床 {len(BM_GEOMETRY_FIELD_MAPPING)} 个')
            
            measured_value_start_idx = 4
            
            logger.info(f'实测值列起始索引: {measured_value_start_idx} (E列)')
            
            available_data_columns = len(detection_header_row) - measured_value_start_idx
            if available_data_columns < len(machine_rows):
                return {
                    'success': False,
                    'message': f'Sheet2数据列数不足：Sheet1有{len(machine_rows)}个机床，但Sheet2从E列开始只有{available_data_columns}列数据。请确保Sheet2的E列、F列等包含足够的数据列。'
                }
            
            logger.info(f'Sheet1机床数量: {len(machine_rows)}, Sheet2可用数据列数: {available_data_columns} (从E列开始)')
        
        # 解析sheet1数据并创建机床记录
        imported_machines = []
        imported_count = 0
        error_count = 0
        success_count = 0
        failure_count = 0
        
        for row_idx, row in enumerate(machine_rows, start=2 if not is_new_template else 3):
            # 每台机床使用独立事务，单台失败不影响其他机床
            try:
                # 跳过空行
                if not any(cell for cell in row):
                    continue
                
                # 根据模板类型提取数据
                if is_new_template:
                    # 新模板列位置
                    # 第0列: 机床类型, 第1列: 机床名称, 第2列: 机床编号, 第3列: 报告编号
                    # 第4列: 规格型号, 第5列: 生产厂家, 第6列: 生产日期
                    
                    # 先读取机床类型，根据类型确定列位置
                    machine_type_raw = row[0] if len(row) > 0 else ''
                    machine_type_map = {
                        'machining_center': 'machining_center',
                        'boring_milling_machine': 'boring_milling_machine',
                        '加工中心': 'machining_center',
                        '镗铣床': 'boring_milling_machine'
                    }
                    machine_type_internal_temp = machine_type_map.get(machine_type_raw, 'machining_center')
                    
                    machine_type = machine_type_raw
                    machine_name = row[1] if len(row) > 1 else ''
                    machine_number = row[2] if len(row) > 2 else None
                    report_number = row[3] if len(row) > 3 else ''
                    model_value = row[4] if len(row) > 4 else None
                    manufacturer = row[5] if len(row) > 5 else None
                    production_date_raw = row[6] if len(row) > 6 else None
                    
                    if machine_type_internal_temp == 'boring_milling_machine':
                        # 镗铣床列位置
                        # 第7列: 实测立柱, 第8列: 实测X, 第9列: 实测Y, 第10列: 实测Z
                        # 第11列: 工作台长边长度, 第12列: 镗轴直径, 第13列: 铣轴直径, 第14列: 工作台类型
                        column_measurement_range = row[7] if len(row) > 7 else None
                        x_axis_travel = row[8] if len(row) > 8 else None
                        y_axis_travel = row[9] if len(row) > 9 else None
                        z_axis_travel = row[10] if len(row) > 10 else None
                        worktable_long_side = row[11] if len(row) > 11 else None
                        boring_axis_diameter = row[12] if len(row) > 12 else None
                        milling_axis_diameter = row[13] if len(row) > 13 else None
                        worktable_type_raw = row[14] if len(row) > 14 else ''
                        # 工作台类型中文到英文映射
                        worktable_type_map = {
                            '固定分度的回转工作台': 'fixed_indexing',
                            '任意角度定位的回转工作台': 'arbitrary_positioning',
                            'fixed_indexing': 'fixed_indexing',
                            'arbitrary_positioning': 'arbitrary_positioning',
                        }
                        worktable_type = worktable_type_map.get(worktable_type_raw, worktable_type_raw) if worktable_type_raw else ''
                        worktable_short_side = None
                        workpiece_size = None
                        spindle_type = ''
                    else:
                        # 加工中心列位置
                        # 第7列: 实测X, 第8列: 实测Y, 第9列: 实测Z
                        # 第10列: 工作台短边长度, 第11列: 工件尺寸, 第12列: 主轴类型, 第13列: 备注
                        x_axis_travel = row[7] if len(row) > 7 else None
                        y_axis_travel = row[8] if len(row) > 8 else None
                        z_axis_travel = row[9] if len(row) > 9 else None
                        worktable_short_side = row[10] if len(row) > 10 else None
                        workpiece_size = row[11] if len(row) > 11 else None
                        spindle_type = row[12] if len(row) > 12 else ''
                        column_measurement_range = None
                        worktable_long_side = None
                        boring_axis_diameter = None
                        milling_axis_diameter = None
                        worktable_type = ''
                else:
                    # 旧模板列位置
                    # 第0列: 序号, 第1列: 机床名称, 第2列: 报告编号, 第3列: 机床类型
                    # 第4列: 测量X, 第5列: 测量Y, 第6列: 测量Z
                    # 第7列: 工作台短边, 第8列: 工件尺寸, 第9列: 主轴类型
                    machine_id = row[0] if len(row) > 0 else None
                    machine_name = row[1] if len(row) > 1 else ''
                    report_number = row[2] if len(row) > 2 else ''
                    machine_type = row[3] if len(row) > 3 else ''
                    x_axis_travel = row[4] if len(row) > 4 else None
                    y_axis_travel = row[5] if len(row) > 5 else None
                    z_axis_travel = row[6] if len(row) > 6 else None
                    worktable_short_side = row[7] if len(row) > 7 else None
                    workpiece_size = row[8] if len(row) > 8 else None
                    spindle_type = row[9] if len(row) > 9 else ''
                    machine_number = row[10] if len(row) > 10 else None
                    model_value = row[11] if len(row) > 11 else None
                    manufacturer = row[12] if len(row) > 12 else None
                    production_date_raw = row[13] if len(row) > 13 else None
                    # 镗铣床特有字段（旧模板不支持）
                    column_measurement_range = None
                    worktable_long_side = None
                    boring_axis_diameter = None
                    milling_axis_diameter = None
                    worktable_type = ''
                
                production_date = parse_year_month(production_date_raw) if production_date_raw else None
                
                # 调试：打印提取的数据
                logger.info(f'第{row_idx}行数据提取结果:')
                logger.info(f'  机床名称: {machine_name}')
                logger.info(f'  报告编号: {report_number}')
                logger.info(f'  机床类型: {machine_type}')
                logger.info(f'  X轴行程: {x_axis_travel}')
                logger.info(f'  Y轴行程: {y_axis_travel}')
                logger.info(f'  Z轴行程: {z_axis_travel}')
                logger.info(f'  生产日期原始值: {production_date_raw}, 解析后: {production_date}')
                if machine_type_internal_temp == 'boring_milling_machine':
                    logger.info(f'  [镗铣床] 立柱测量范围: {column_measurement_range}')
                    logger.info(f'  [镗铣床] 工作台X轴: {x_axis_travel}')
                    logger.info(f'  [镗铣床] 主轴箱Y轴: {y_axis_travel}')
                    logger.info(f'  [镗铣床] Z轴测量范围: {z_axis_travel}')
                    logger.info(f'  [镗铣床] 工作台长边: {worktable_long_side}')
                    logger.info(f'  [镗铣床] 镗轴直径: {boring_axis_diameter}')
                    logger.info(f'  [镗铣床] 铣轴直径: {milling_axis_diameter}')
                    logger.info(f'  [镗铣床] 工作台类型: {worktable_type}')
                
                # 转换机床类型为系统内部值
                machine_type_map = {
                    'machining_center': 'machining_center',
                    'boring_milling_machine': 'boring_milling_machine',
                    '加工中心': 'machining_center',
                    '镗铣床': 'boring_milling_machine'
                }
                machine_type_internal = machine_type_map.get(machine_type, 'machining_center')
                
                # 数据兼容性处理：主轴类型
                spindle_type_map = {
                    '整体主轴': 'integral',
                    '附加主轴': 'attached',

                }
                if spindle_type:
                    spindle_type_internal = spindle_type_map.get(spindle_type, spindle_type)
                else:
                    spindle_type_internal = 'integral'
                
                # 数据兼容性处理：工件尺寸
                if workpiece_size:
                    if isinstance(workpiece_size, str) and 'mm' in workpiece_size:
                        workpiece_size = extract_number(workpiece_size)
                    elif str(workpiece_size) in ['80', '160', '320']:
                        workpiece_size = int(workpiece_size)
                    else:
                        workpiece_size = extract_number(workpiece_size)
                
                # 检查用户是否为管理员
                try:
                    user_profile = request.user.userprofile
                    is_admin = user_profile.is_admin()
                except Exception:
                    is_admin = request.user.is_superuser

                # 使用事务保存机床数据，确保原子性
                with transaction.atomic():
                    # 检查是否已存在相同名称、报告编号且属于当前用户的机床
                    # 这样不同用户可以各自拥有自己的机床副本
                    existing_machine = Machine.objects.filter(
                        name=machine_name, 
                        report_number=report_number,
                        created_by=request.user
                    ).first()
                    
                    if existing_machine:
                        # 更新现有记录
                        logger.info(f'更新现有机床: {machine_name} (报告编号: {report_number}, ID: {existing_machine.id})')
                        machine = existing_machine
                        machine.model = machine_name
                        machine.machine_type = machine_type_internal
                        machine.x_axis_travel = extract_number(x_axis_travel)
                        machine.y_axis_travel = extract_number(y_axis_travel)
                        machine.z_axis_travel = extract_number(z_axis_travel)
                        machine.worktable_length = extract_number(worktable_short_side)
                        machine.workpiece_size = workpiece_size
                        machine.spindle_type = spindle_type_internal
                        if machine_number:
                            machine.machine_number = machine_number
                        if model_value:
                            machine.model = model_value
                        if manufacturer:
                            machine.manufacturer = manufacturer
                        if production_date:
                            machine.production_date = production_date
                        machine.save()
                    else:
                        # 创建新记录
                        machine = Machine(
                            name=machine_name,
                            model=model_value,
                            report_number=report_number,
                            machine_number=machine_number,
                            manufacturer=manufacturer,
                            production_date=production_date,
                            machine_type=machine_type_internal,
                            x_axis_travel=extract_number(x_axis_travel),
                            y_axis_travel=extract_number(y_axis_travel),
                            z_axis_travel=extract_number(z_axis_travel),
                            worktable_length=extract_number(worktable_short_side),
                            workpiece_size=workpiece_size,
                            spindle_type=spindle_type_internal,
                            created_by=request.user,
                            admin_saved=is_admin
                        )

                    # 保存机床记录
                    machine.save()

                    # 镗铣床特有字段也保存到Django模型
                    if machine_type_internal == 'boring_milling_machine':
                        machine.worktable_length_lc = extract_number(worktable_long_side)
                        machine.column_measurement_range = extract_number(column_measurement_range)
                        machine.worktable_x_range = extract_number(x_axis_travel)
                        machine.spindle_box_y_range = extract_number(y_axis_travel)
                        machine.z_axis_measurement_range = extract_number(z_axis_travel)
                        machine.boring_axis_diameter = extract_number(boring_axis_diameter)
                        machine.milling_axis_diameter = extract_number(milling_axis_diameter)
                        if worktable_type:
                            machine.worktable_type = worktable_type
                        machine.save()

                    logger.info(f'成功保存机床记录: {machine_name} (ID: {machine.id})')

                    # 添加到导入列表
                    imported_machines.append({
                        'id': machine.id,
                        'name': machine_name,
                        'report_number': report_number,
                        'machine_type': machine_type_internal
                    })

                    # 自动保存到达梦数据库
                    if machine_type_internal == 'boring_milling_machine':
                        # 镗铣床特有字段
                        dameng_data = {
                            'column_measurement_range': extract_number(column_measurement_range),
                            'worktable_x_range': extract_number(x_axis_travel),
                            'spindle_box_y_range': extract_number(y_axis_travel),
                            'z_axis_measurement_range': extract_number(z_axis_travel),
                            'worktable_length_lc': extract_number(worktable_long_side),
                            'boring_axis_diameter': extract_number(boring_axis_diameter),
                            'milling_axis_diameter': extract_number(milling_axis_diameter),
                            'worktable_type': worktable_type,
                        }
                        logger.info(f'镗铣床达梦数据库保存数据: {dameng_data}')
                        result = DamengService.update_boring_milling_parameters(machine.id, dameng_data, request.user.username, request.META.get('REMOTE_ADDR', '127.0.0.1'))
                        if result['success']:
                            logger.info(f'成功保存到达梦数据库（镗铣床）: {machine_name} (ID: {machine.id})')
                        else:
                            logger.warning(f'保存到达梦数据库失败: {result["message"]}')
                    else:
                        dameng_data = {
                            'machine_type': machine_type_internal,
                            'x_axis_travel': extract_number(x_axis_travel),
                            'y_axis_travel': extract_number(y_axis_travel),
                            'z_axis_travel': extract_number(z_axis_travel),
                            'worktable_length': extract_number(worktable_short_side),
                            'workpiece_size': workpiece_size,
                            'spindle_type': spindle_type_internal
                        }
                        result = DamengService.update_machine_parameters(machine.id, dameng_data, request.user.username, request.META.get('REMOTE_ADDR', '127.0.0.1'))
                        if result['success']:
                            logger.info(f'成功保存到达梦数据库: {machine_name} (ID: {machine.id})')
                        else:
                            logger.warning(f'保存到达梦数据库失败: {result["message"]}')

                    # 计算并保存允差值
                    try:
                        from apps.evaluation.views import calculate_tolerances_for_group_1
                        tolerance_group_1 = calculate_tolerances_for_group_1(
                            machine_type_internal,
                            extract_number(x_axis_travel),
                            extract_number(y_axis_travel),
                            extract_number(z_axis_travel),
                            extract_number(worktable_short_side),
                            workpiece_size,
                            spindle_type_internal
                        )
                        machine.tolerance_group_1 = tolerance_group_1
                        machine.save()
                        logger.info(f'成功计算并保存允差值: {machine_name} (ID: {machine.id})')
                    except Exception as e:
                        logger.error(f'计算允差值时出错: {str(e)}')
                
                imported_count += 1
                
            except Exception as e:
                error_count += 1
                error_messages.append(f"第{row_idx}行：{str(e)}")
                logger.error(f'处理第{row_idx}行时出错: {str(e)}')
        
        # 处理sheet2数据（检测项目数据）- 仅对旧模板
        if imported_count > 0 and not is_new_template:
            logger.info('开始处理检测项目数据（sheet2）')
            logger.info(f'导入的机床数量: {imported_count}')
            logger.info(f'导入的机床列表: {imported_machines}')
            
            success_count = 0
            failure_count = 0
            detection_errors = []
            
            try:
                
                for machine_idx, machine_info in enumerate(imported_machines):
                    machine_id = machine_info['id']
                    machine_name = machine_info['name']
                    is_boring_milling = machine_info.get('machine_type') == 'boring_milling_machine'
                    current_field_mapping = BM_GEOMETRY_FIELD_MAPPING if is_boring_milling else GEOMETRY_FIELD_MAPPING
                    
                    measured_value_col_idx = measured_value_start_idx + machine_idx
                    
                    logger.info(f'处理机床 {machine_name} (ID: {machine_id}) 的检测项目数据（第{measured_value_col_idx + 1}列，即{chr(65 + measured_value_col_idx)}列），机床类型: {"镗铣床" if is_boring_milling else "加工中心"}')
                    
                    if measured_value_col_idx >= len(detection_header_row):
                        error_msg = f'机床 {machine_name} (ID: {machine_id}) 的数据列索引{measured_value_col_idx}超出范围'
                        logger.error(error_msg)
                        detection_errors.append(error_msg)
                        failure_count += 1
                        continue
                    
                    detection_data = {}
                    
                    for row_idx, row in enumerate(detection_rows, start=2):
                        if not any(cell for cell in row):
                            continue
                        
                        item_description = row[2] if len(row) > 2 else None
                        
                        if not item_description:
                            continue
                        
                        item_description = str(item_description).strip()
                        
                        field_name = current_field_mapping.get(item_description)
                        
                        if not field_name:
                            logger.debug(f'未找到映射: {item_description}')
                            continue
                        
                        if measured_value_col_idx < len(row):
                            value = row[measured_value_col_idx]
                            logger.debug(f'机床{machine_name} - 行{row_idx} - 列{measured_value_col_idx}({chr(65 + measured_value_col_idx)}列) - {item_description} - 原始值: {value}')
                            
                            numeric_value = extract_number(value)
                            
                            if numeric_value is not None:
                                detection_data[field_name] = numeric_value
                                logger.info(f'{item_description} -> {field_name} = {numeric_value}')
                            else:
                                logger.warning(f'{item_description} 原始值 {value} 无法转换为数字')
                        else:
                            logger.warning(f'{item_description} 列索引{measured_value_col_idx}超出范围，行长度: {len(row)}')
                    
                    logger.info(f'成功解析 {len(detection_data)} 个检测项目')
                    
                    if len(detection_data) > 0:
                        logger.info(f'机床 {machine_name} (ID: {machine_id}) 的检测项目数据:')
                        for field, value in detection_data.items():
                            logger.info(f'  {field}: {value}')
                    else:
                        logger.warning(f'机床 {machine_name} (ID: {machine_id}) 没有解析到任何检测项目数据')
                    
                    expected_fields = len(current_field_mapping)
                    if len(detection_data) < expected_fields * 0.5:
                        warning_msg = f'机床 {machine_name} (ID: {machine_id}) 的检测项目数据不完整，仅{len(detection_data)}/{expected_fields}个字段有数据'
                        logger.warning(warning_msg)
                        detection_errors.append(warning_msg)
                    
                    # 保存检测项目数据到达梦数据库
                    if detection_data:
                        try:
                            logger.info(f'开始保存机床 {machine_name} (ID: {machine_id}) 的检测项目数据到达梦数据库')
                            
                            if machine_info.get('machine_type') == 'boring_milling_machine':
                                result = DamengService.update_boring_milling_parameters(
                                    machine_id=machine_id,
                                    update_data=detection_data,
                                    user=request.user.username,
                                    ip_address=request.META.get('REMOTE_ADDR', '127.0.0.1')
                                )
                            else:
                                result = DamengService.update_machine_parameters(
                                    machine_id=machine_id,
                                    update_data=detection_data,
                                    user=request.user.username,
                                    ip_address=request.META.get('REMOTE_ADDR', '127.0.0.1')
                                )
                            
                            if result['success']:
                                logger.info(f'成功保存检测项目数据到达梦数据库: 机床 {machine_name} (ID: {machine_id})')
                                success_count += 1
                            else:
                                error_msg = f'保存机床 {machine_name} (ID: {machine_id}) 的检测项目数据失败: {result["message"]}'
                                logger.warning(error_msg)
                                detection_errors.append(error_msg)
                                failure_count += 1
                        except Exception as e:
                            error_msg = f'保存机床 {machine_name} (ID: {machine_id}) 的检测项目数据时出错: {str(e)}'
                            logger.error(error_msg)
                            import traceback
                            logger.error(traceback.format_exc())
                            detection_errors.append(error_msg)
                            failure_count += 1
                    else:
                        error_msg = f'机床 {machine_name} (ID: {machine_id}) 没有检测项目数据'
                        logger.warning(error_msg)
                        detection_errors.append(error_msg)
                        failure_count += 1
                
                logger.info(f'检测项目数据处理完成: 成功 {success_count}, 失败 {failure_count}')
                
                if detection_errors:
                    error_messages.extend(detection_errors)
                
            except Exception as e:
                error_msg = f'处理检测项目数据时出错: {str(e)}'
                logger.error(error_msg)
                import traceback
                logger.error(traceback.format_exc())
                error_messages.append(error_msg)
        
        module_success_count = 0
        module_failure_count = 0
        module_errors = []
        
        for sheet_name in input_workbook.sheetnames:
            # 跳过机床信息工作表和旧模板的Sheet1、Sheet2
            if sheet_name == '机床信息' or sheet_name == 'Sheet1' or sheet_name == 'Sheet2':
                continue
            
            logger.info(f'检查工作表: {sheet_name}')
            
            try:
                sheet = input_workbook[sheet_name]
                
                if sheet_name.startswith('1.1') or ('几何精度' in sheet_name and '保持性' not in sheet_name):
                    logger.info(f'匹配到几何精度工作表: {sheet_name}')
                    result = _parse_and_save_geometry_precision(sheet, imported_machines, DamengService, request)
                elif sheet_name.startswith('1.2') or ('位置精度' in sheet_name and '保持性' not in sheet_name):
                    logger.info(f'匹配到位置精度工作表: {sheet_name}')
                    result = _parse_and_save_position_precision(sheet, imported_machines, DamengService, request)
                elif sheet_name.startswith('1.3') or ('加工精度' in sheet_name and '保持性' not in sheet_name):
                    logger.info(f'匹配到加工精度工作表: {sheet_name}')
                    result = _parse_and_save_machining_precision(sheet, imported_machines, DamengService, request)
                elif sheet_name.startswith('2.1') or ('加工性能' in sheet_name and '刚度' not in sheet_name):
                    logger.info(f'匹配到加工性能工作表: {sheet_name}')
                    result = _parse_and_save_machining_performance(sheet, imported_machines, DamengService, request)
                elif sheet_name.startswith('2.2') or '刚度' in sheet_name:
                    logger.info(f'匹配到刚度工作表: {sheet_name}')
                    result = _parse_and_save_stiffness(sheet, imported_machines, DamengService, request)
                elif sheet_name.startswith('2.3') or '运行性能' in sheet_name:
                    logger.info(f'匹配到运行性能工作表: {sheet_name}')
                    result = _parse_and_save_running_performance(sheet, imported_machines, DamengService, request)
                elif sheet_name.startswith('2.4') or '负荷' in sheet_name:
                    logger.info(f'匹配到负荷能力工作表: {sheet_name}')
                    result = _parse_and_save_load_capacity(sheet, imported_machines, DamengService, request)
                elif sheet_name.startswith('3.1') or '故障间隔' in sheet_name:
                    logger.info(f'匹配到MTBF工作表: {sheet_name}')
                    result = _parse_and_save_mtbf(sheet, imported_machines, DamengService, request)
                elif sheet_name.startswith('3.2') or '修复时间' in sheet_name:
                    logger.info(f'匹配到修复时间工作表: {sheet_name}')
                    result = _parse_and_save_repair_time(sheet, imported_machines, DamengService, request)
                elif sheet_name.startswith('3.3') or ('故障率' in sheet_name and '首次' not in sheet_name):
                    logger.info(f'匹配到故障率工作表: {sheet_name}')
                    result = _parse_and_save_failure_rate(sheet, imported_machines, DamengService, request)
                elif sheet_name.startswith('3.4') or '一致性' in sheet_name:
                    logger.info(f'匹配到一致性工作表: {sheet_name}')
                    result = _parse_and_save_consistency(sheet, imported_machines, DamengService, request)
                elif sheet_name.startswith('3.5') or '保持性' in sheet_name:
                    logger.info(f'匹配到精度保持性工作表: {sheet_name}')
                    result = _parse_and_save_precision_retention(sheet, imported_machines, DamengService, request)
                elif sheet_name.startswith('3.6') or '首次故障' in sheet_name:
                    logger.info(f'匹配到MTTF工作表: {sheet_name}')
                    result = _parse_and_save_mttf(sheet, imported_machines, DamengService, request)
                elif sheet_name.startswith('4') or '智能' in sheet_name:
                    logger.info(f'匹配到智能化工作表: {sheet_name}')
                    result = _parse_and_save_intelligence(sheet, imported_machines, DamengService, request)
                elif sheet_name.startswith('5.1') or '资源' in sheet_name or '耗材' in sheet_name or '切削液' in sheet_name or '润滑油' in sheet_name:
                    logger.info(f'匹配到资源循环与耗材工作表: {sheet_name}')
                    result = _parse_and_save_resource(sheet, imported_machines, DamengService, request)
                elif sheet_name.startswith('5.2') or '能源' in sheet_name or '能效' in sheet_name or '待机功率' in sheet_name or '空运转功率' in sheet_name:
                    logger.info(f'匹配到能源-能效工作表: {sheet_name}')
                    result = _parse_and_save_energy(sheet, imported_machines, DamengService, request)
                else:
                    logger.info(f'工作表 {sheet_name} 未匹配到任何模块，跳过')
                    continue
                
                if result.get('success'):
                    module_success_count += result.get('count', 1)
                else:
                    module_failure_count += 1
                    if result.get('error'):
                        module_errors.append(result['error'])
            
            except Exception as e:
                logger.error(f'处理工作表 {sheet_name} 出错: {str(e)}')
                module_errors.append(f'{sheet_name}处理错误: {str(e)}')
        
        os.unlink(temp_file_path)
        
        message_parts = []
        message_parts.append(f'成功导入 {imported_count} 条机床记录')
        
        if success_count > 0:
            message_parts.append(f'成功保存 {success_count} 条检测项目数据')
        
        if failure_count > 0:
            message_parts.append(f'{failure_count} 条检测项目数据保存失败')
        
        if module_success_count > 0:
            message_parts.append(f'模块数据成功 {module_success_count} 项')
        
        if module_failure_count > 0:
            message_parts.append(f'模块数据失败 {module_failure_count} 项')
        
        if error_messages:
            message_parts.append(f'共 {len(error_messages)} 个警告或错误')
        
        if module_errors:
            message_parts.append(f'模块错误: {len(module_errors)} 个')
        
        final_message = '；'.join(message_parts)
        
        logger.info(f'Excel文件导入完成: {final_message}')
        
        return {
            'success': True,
            'message': final_message,
            'imported_count': imported_count,
            'detection_success_count': success_count,
            'detection_failure_count': failure_count,
            'module_success_count': module_success_count,
            'module_failure_count': module_failure_count,
            'error_count': error_count,
            'error_messages': error_messages + module_errors
        }
    
    except Exception as e:
        logger.error(f'处理Excel文件时出错: {str(e)}')
        import traceback
        logger.error(traceback.format_exc())
        return {'success': False, 'message': f'无法识别该Excel文件: {str(e)}'}


def _parse_and_save_geometry_precision(sheet, imported_machines, DamengService, request):
    """
    解析并保存几何精度数据
    
    Args:
        sheet: Excel工作表对象
        imported_machines: 已导入的机床列表
        DamengService: 达梦数据库服务
        request: HTTP请求对象
        
    Returns:
        包含success和count的字典
    """
    success_count = 0
    for machine_idx, machine_info in enumerate(imported_machines):
        data = {}
        is_boring_milling = machine_info.get('machine_type') == 'boring_milling_machine'
        current_mapping = BM_GEOMETRY_FIELD_MAPPING if is_boring_milling else GEOMETRY_FIELD_MAPPING
        
        measured_value_col_idx = 4 + machine_idx
        
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < measured_value_col_idx + 1:
                continue
            description = str(row[2]).strip() if row[2] else ''
            measured_value = row[measured_value_col_idx] if measured_value_col_idx < len(row) else None
            
            if description in current_mapping and measured_value is not None:
                data[current_mapping[description]] = extract_number(measured_value)
        
        if data:
            result = update_machine_parameters_by_type(DamengService, machine_info, data, request)
            if result['success']:
                logger.info(f'成功保存几何精度数据: {machine_info["name"]}')
                success_count += 1
    
    return {'success': True, 'count': success_count}


def _parse_and_save_position_precision(sheet, imported_machines, DamengService, request):
    """
    解析并保存位置精度数据
    
    包括线性轴和回转轴的重复定位精度、定位精度、位置偏差等参数
    """
    linear_axis_mapping = {
        '双向重复定位精度R': 'linear_axis_repeatability',
        '双向定位精度A': 'linear_axis_positioning_accuracy',
        '轴线的平均双向位置偏差范围M': 'linear_axis_mean_deviation',
        '双向定位系统偏差E': 'linear_axis_system_deviation',
        '轴线的反向差值B': 'linear_axis_backlash',
    }
    
    rotary_axis_mapping = {
        '双向重复定位精度R': 'rotary_axis_repeatability',
        '双向定位精度A': 'rotary_axis_positioning_accuracy',
        '轴线的平均双向位置偏差范围M': 'rotary_axis_mean_deviation',
        '双向定位系统偏差E': 'rotary_axis_system_deviation',
        '轴线的反向差值B': 'rotary_axis_backlash',
    }
    
    bm_linear_axis_mapping = {
        '双向重复定位精度R': 'bm_linear_axis_repeatability',
        '双向定位精度A': 'bm_linear_axis_positioning_accuracy',
        '轴线的平均双向位置偏差范围M': 'bm_linear_axis_mean_deviation',
        '双向定位系统偏差E': 'bm_linear_axis_system_deviation',
        '轴线的反向差值B': 'bm_linear_axis_backlash',
    }
    
    bm_rotary_axis_mapping = {
        '双向重复定位精度R': 'bm_rotary_axis_repeatability',
        '双向定位精度A': 'bm_rotary_axis_positioning_accuracy',
        '轴线的平均双向位置偏差范围M': 'bm_rotary_axis_mean_deviation',
        '双向定位系统偏差E': 'bm_rotary_axis_system_deviation',
        '轴线的反向差值B': 'bm_rotary_axis_backlash',
    }
    
    axis_suffixes = ['_x', '_y', '_z']
    rotary_suffixes = ['_a', '_b', '_c']
    bm_axis_suffixes = ['_x', '_y', '_z', '_w']
    
    success_count = 0
    for machine_idx, machine_info in enumerate(imported_machines):
        data = {}
        current_section = None
        is_boring_milling = machine_info.get('machine_type') == 'boring_milling_machine'
        
        if is_boring_milling:
            current_linear_mapping = bm_linear_axis_mapping
            current_rotary_mapping = bm_rotary_axis_mapping
            current_axis_suffixes = bm_axis_suffixes
        else:
            current_linear_mapping = linear_axis_mapping
            current_rotary_mapping = rotary_axis_mapping
            current_axis_suffixes = axis_suffixes
        
        base_col_idx = 3 + machine_idx * (4 if is_boring_milling else 3)
        
        for row_idx, row in enumerate(sheet.iter_rows(min_row=1, values_only=True), 1):
            if not row or len(row) < base_col_idx + len(current_axis_suffixes):
                continue
            
            first_col = str(row[0]).strip() if row[0] else ''
            description = str(row[2]).strip() if row[2] else ''
            
            if first_col == '线性轴':
                current_section = 'linear'
                continue
            elif first_col == '回转轴':
                current_section = 'rotary'
                continue
            
            if description in ['描述', '项目编号', '子项'] or not description:
                continue
            
            if current_section == 'linear' and description in current_linear_mapping:
                base_field = current_linear_mapping[description]
                for axis_idx, suffix in enumerate(current_axis_suffixes):
                    col_idx = base_col_idx + axis_idx
                    value = row[col_idx] if col_idx < len(row) else None
                    if value is not None:
                        data[base_field + suffix] = extract_number(value)
                        
            elif current_section == 'rotary' and description in current_rotary_mapping:
                base_field = current_rotary_mapping[description]
                for axis_idx, suffix in enumerate(rotary_suffixes):
                    col_idx = base_col_idx + axis_idx
                    value = row[col_idx] if col_idx < len(row) else None
                    if value is not None:
                        data[base_field + suffix] = extract_number(value)
        
        if data:
            result = update_machine_parameters_by_type(DamengService, machine_info, data, request)
            if result['success']:
                logger.info(f'成功保存位置精度数据: {machine_info["name"]}, 数据: {data}')
                success_count += 1
    
    return {'success': True, 'count': success_count}


def _parse_and_save_machining_precision(sheet, imported_machines, DamengService, request):
    """
    解析并保存加工精度数据
    
    包括圆柱度、垂直度、平行度、圆度等加工精度参数
    镗铣床支持正方形、菱形、圆、镗孔等不同加工项目
    """
    field_mapping = {
        '圆柱度': 'cylindricity',
        '孔中心轴线与基孔A的垂直度': 'hole_perpendicularity_a',
        '相邻面与基面B的垂直度': 'adjacent_perpendicularity_b',
        '相对面与基面B的平行度': 'opposite_parallelism_b',
        '侧面对基准B的倾斜度': 'side_inclination_b',
        '圆度': 'circularity',
        '外圆和内圆孔C的同心度': 'concentricity_c',
        '斜面的直线度': 'slope_linearity',
        '3°角斜面对基准B的倾斜度': 'slope_inclination_b',
        '斜面对基准B的倾斜度': 'slope_inclination_b',
        '孔相对于内孔C的位置度': 'hole_position_c',
        '内孔与外孔D的同心度': 'concentricity_d',
    }
    
    bm_machining_mapping = {
        '正方形': {
            '侧面的直线度': 'bm_square_a',
            '相邻面对基准B的垂直度': 'bm_square_b',
            '相对面对基准B的平行度': 'bm_square_c',
            '侧面对基准B的垂直度': 'bm_square_d',
            '相邻面对基准B的平行度': 'bm_square_e',
            '相对面对基准B的垂直度': 'bm_square_f',
        },
        '菱形': {
            '侧面的直线度': 'bm_diamond_d',
            '侧面对基面B的倾斜度': 'bm_diamond_e',
            '侧面对基准B的倾斜度': 'bm_diamond_e',
        },
        '圆': {
            '圆度': 'bm_circle_f',
            '外圆和内圆孔C的同心度': 'bm_circle_g',
            '同心度': 'bm_circle_g',
        },
        '镗孔': {
            '孔对基准A的位置度': 'bm_boring_hole_h',
            '孔对基准A的同心度': 'bm_boring_hole_i',
        },
        '刀库': {
            '刀库定位精度': 'bm_tool_magazine_positioning_accuracy',
            '刀库重复定位精度': 'bm_tool_magazine_repeatability',
            '定位精度': 'bm_tool_magazine_positioning_accuracy',
            '重复定位精度': 'bm_tool_magazine_repeatability',
        },
    }
    
    success_count = 0
    for machine_idx, machine_info in enumerate(imported_machines):
        data = {}
        is_boring_milling = machine_info.get('machine_type') == 'boring_milling_machine'
        row_index = 0
        current_item = None
        
        # 镗铣床新模板：描述在第4列(索引3)，实测值从第5列(索引4)开始
        # 加工中心旧模板：描述在第3列(索引2)，实测值从第4列(索引3)开始
        if is_boring_milling:
            description_col_idx = 3
            measured_value_col_idx = 4 + machine_idx
        else:
            description_col_idx = 2
            measured_value_col_idx = 3 + machine_idx
        
        for row in sheet.iter_rows(min_row=2, values_only=True):
            row_index += 1
            if not row or len(row) < measured_value_col_idx + 1:
                continue
            
            item_name = str(row[1]).strip() if row[1] else None
            description = str(row[description_col_idx]).strip() if len(row) > description_col_idx and row[description_col_idx] else ''
            measured_value = row[measured_value_col_idx] if measured_value_col_idx < len(row) else None
            
            if is_boring_milling:
                if item_name:
                    current_item = item_name
                
                if current_item and current_item in bm_machining_mapping and description:
                    item_mapping = bm_machining_mapping[current_item]
                    if description in item_mapping:
                        if measured_value is not None:
                            data[item_mapping[description]] = extract_number(measured_value)
            else:
                if description in field_mapping:
                    if measured_value is not None:
                        data[field_mapping[description]] = extract_number(measured_value)
                    else:
                        data[field_mapping[description]] = 0
                elif description == '侧面的直线度':
                    if row_index == 3:
                        data['side_linearity'] = extract_number(measured_value) if measured_value is not None else 0
                    elif row_index == 6:
                        data['side_linearity_2'] = extract_number(measured_value) if measured_value is not None else 0
        
        if data:
            result = update_machine_parameters_by_type(DamengService, machine_info, data, request)
            if result['success']:
                logger.info(f'成功保存加工精度数据: {machine_info["name"]}, 数据: {data}')
                success_count += 1
    
    return {'success': True, 'count': success_count}


def _parse_and_save_machining_performance(sheet, imported_machines, DamengService, request):
    """
    解析并保存加工性能数据
    
    包括CP值、转速、速度、定位精度等加工性能参数
    """
    success_count = 0
    for machine_idx, machine_info in enumerate(imported_machines):
        data = {}
        # 第一个机床的数据在第2列（索引1），第二个机床在第3列（索引2），以此类推
        value_col_idx = 1 + machine_idx
        
        for row in sheet.iter_rows(min_row=1, values_only=True):
            if not row or len(row) < value_col_idx + 1:
                continue
            param_name = str(row[0]).strip() if row[0] else ''
            value = row[value_col_idx] if value_col_idx < len(row) else None
            
            if param_name == 'CP值':
                data['cp_value'] = extract_number(value) if value else 0
            elif '峰值转速' in param_name or 'N_peak' in param_name:
                data['peak_speed'] = extract_number(value) if value else 0
            elif '谷值转速' in param_name or 'N_valley' in param_name:
                data['valley_speed'] = extract_number(value) if value else 0
            elif '设计最大转速' in param_name or 'N_design_max' in param_name:
                data['design_max_speed'] = extract_number(value) if value else 0
            elif '最高快移速度' in param_name:
                data['max_fast_speed'] = extract_number(value) if value else 0
            elif '设计最高速度' in param_name or 'V_design_max' in param_name:
                data['design_max_fast_speed'] = extract_number(value) if value else 0
            elif '最高速度下' in param_name and '重复定位精度' in param_name:
                data['repeat_positioning_accuracy'] = extract_number(value) if value else 0
            elif '最大进给速度' in param_name and '重复定位精度' not in param_name:
                data['max_feed_speed'] = extract_number(value) if value else 0
            elif '目标最高速度' in param_name:
                data['target_max_speed'] = extract_number(value) if value else 0
            elif ('R2' in param_name or '及加速度下' in param_name) and '重复定位精度' in param_name:
                data['repeat_positioning_accuracy_feed'] = extract_number(value) if value else 0
        
        if data:
            result = update_machine_parameters_by_type(DamengService, machine_info, data, request)
            if result['success']:
                logger.info(f'成功保存加工性能数据: {machine_info["name"]}, 数据: {data}')
                success_count += 1
    
    return {'success': True, 'count': success_count}


def _parse_and_save_stiffness(sheet, imported_machines, DamengService, request):
    """
    解析并保存机床刚度数据
    
    包括径向刚度、轴向刚度、切削刚度、静刚度等参数
    """
    success_count = 0
    for machine_idx, machine_info in enumerate(imported_machines):
        data = {}
        # 第一个机床的数据在第2列（索引1），第二个机床在第3列（索引2），以此类推
        value_col_idx = 1 + machine_idx
        
        for row in sheet.iter_rows(min_row=1, values_only=True):
            if not row or len(row) < value_col_idx + 1:
                continue
            param_name = str(row[0]).strip() if row[0] else ''
            value = row[value_col_idx] if value_col_idx < len(row) else None
            
            if '径向刚度' in param_name:
                data['radial_stiffness'] = extract_number(value)
            elif '轴向刚度' in param_name:
                data['axial_stiffness'] = extract_number(value)
            elif param_name == '机床类型':
                machine_type_value = str(value).strip() if value else 'tool-rotating'
                if '刀具旋转' in machine_type_value or '加工中心' in machine_type_value or '铣' in machine_type_value:
                    data['radial_machine_type'] = 'tool-rotating'
                    data['axial_machine_type'] = 'tool-rotating'
                elif '工件旋转' in machine_type_value or '车' in machine_type_value:
                    data['radial_machine_type'] = 'workpiece-rotating'
                    data['axial_machine_type'] = 'workpiece-rotating'
                else:
                    data['radial_machine_type'] = 'tool-rotating'
                    data['axial_machine_type'] = 'tool-rotating'
            elif '切削刚度' in param_name:
                data['cutting_stiffness'] = extract_number(value)
            elif '合格线' in param_name:
                data['static_stiffness_accept'] = extract_number(value)
            elif '静刚度 K_s' in param_name or param_name == '静刚度 K_s':
                data['static_stiffness_value'] = extract_number(value)
        
        if data:
            result = update_machine_parameters_by_type(DamengService, machine_info, data, request)
            if result['success']:
                logger.info(f'成功保存机床刚度数据: {machine_info["name"]}, 数据: {data}')
                success_count += 1
    
    return {'success': True, 'count': success_count}


def _parse_and_save_running_performance(sheet, imported_machines, DamengService, request):
    """
    解析并保存运行性能数据
    
    包括振动速度、噪声、功率、电流波动、热位移等运行性能参数
    """
    success_count = 0
    logger.info(f'[DEBUG] _parse_and_save_running_performance: imported_machines数量={len(imported_machines)}')
    
    for machine_idx, machine_info in enumerate(imported_machines):
        data = {}
        value_col_idx = 1 + machine_idx
        logger.info(f'[DEBUG] 处理机床: machine_idx={machine_idx}, machine_id={machine_info["id"]}, machine_name={machine_info["name"]}, value_col_idx={value_col_idx}')
        
        for row in sheet.iter_rows(min_row=1, values_only=True):
            if not row or len(row) < value_col_idx + 1:
                continue
            param_name = str(row[0]).strip() if row[0] else ''
            value = row[value_col_idx] if value_col_idx < len(row) else None
            
            if '电压波动' in param_name or '频率波动' in param_name:
                logger.info(f'[DEBUG] 运行性能参数: param_name={param_name}, value={value}, value_type={type(value)}, value_col_idx={value_col_idx}')
            
            if '最大振动速度' in param_name:
                data['max_vibration_speed'] = extract_number(value) if value else 0
            elif '噪声值' in param_name:
                data['noise_value'] = extract_number(value) if value else 0
            elif param_name == '机床类型':
                data['noise_machine_type'] = str(value).strip() if value else None
            elif '重量选择' in param_name:
                data['noise_weight'] = str(value).strip() if value else None
            elif '初检功率' in param_name:
                data['initial_power'] = extract_number(value) if value else 0
            elif '终检功率' in param_name:
                data['final_power'] = extract_number(value) if value else 0
            elif '电流不平衡量' in param_name:
                data['current_unbalance'] = extract_number(value) if value else 0
            elif '电流波动度' in param_name:
                data['current_fluctuation'] = extract_number(value) if value else 0
            elif '电压波动' in param_name:
                logger.info(f'[DEBUG] 电压波动参数识别: param_name={param_name}, value={value}')
                min_val, max_val = extract_range(value)
                logger.info(f'[DEBUG] 电压波动解析结果: min_val={min_val}, max_val={max_val}')
                data['voltage_fluctuation_min'] = min_val if min_val is not None else 0
                data['voltage_fluctuation_max'] = max_val if max_val is not None else 0
            elif '频率波动' in param_name:
                logger.info(f'[DEBUG] 频率波动参数识别: param_name={param_name}, value={value}')
                min_val, max_val = extract_range(value)
                logger.info(f'[DEBUG] 频率波动解析结果: min_val={min_val}, max_val={max_val}')
                data['frequency_fluctuation_min'] = min_val if min_val is not None else 0
                data['frequency_fluctuation_max'] = max_val if max_val is not None else 0
            elif '径向热位移' in param_name:
                data['radial_thermal_displacement'] = extract_number(value) if value else 0
            elif '热平衡梯度 G1' in param_name or ('热平衡' in param_name and '径向' not in param_name and 'G1' in param_name):
                data['radial_thermal_gradient'] = extract_number(value) if value else 0
            elif '轴向热位移' in param_name:
                data['axial_thermal_displacement'] = extract_number(value) if value else 0
            elif '热平衡梯度 G2' in param_name or ('热平衡' in param_name and '轴向' not in param_name and 'G2' in param_name):
                data['axial_thermal_gradient'] = extract_number(value) if value else 0
        
        if data:
            result = update_machine_parameters_by_type(DamengService, machine_info, data, request)
            if result['success']:
                logger.info(f'成功保存运行性能数据: {machine_info["name"]}, 数据: {data}')
                success_count += 1
    
    return {'success': True, 'count': success_count}


def _parse_and_save_load_capacity(sheet, imported_machines, DamengService, request):
    """
    解析并保存负荷能力数据
    
    包括主轴功率、扭矩、负重、切削参数等负荷能力参数
    """
    success_count = 0
    for machine_idx, machine_info in enumerate(imported_machines):
        data = {}
        vibration_count = 0
        temp_machine_type = None
        temp_evaluation_scheme = None
        # 第一个机床的数据在第2列（索引1），第二个机床在第3列（索引2），以此类推
        value_col_idx = 1 + machine_idx
        
        for row in sheet.iter_rows(min_row=1, values_only=True):
            if not row or len(row) < value_col_idx + 1:
                continue
            param_name = str(row[0]).strip() if row[0] else ''
            value = row[value_col_idx] if value_col_idx < len(row) else None
            
            if '实测主轴最大功率' in param_name:
                data['measured_max_power'] = extract_number(value) if value else 0
            elif '设计主轴最大功率' in param_name or '设计最大功率' in param_name:
                data['design_max_power'] = extract_number(value) if value else 0
            elif '切削振动速度' in param_name:
                vibration_count += 1
                data[f'cutting_vibration_speed{vibration_count}'] = extract_number(value) if value else 0
            elif '实测主轴最大扭矩' in param_name or '实测最大扭矩' in param_name:
                data['measured_max_torque'] = extract_number(value) if value else 0
            elif '设计主轴最大扭矩' in param_name or '设计最大扭矩' in param_name:
                data['design_max_torque'] = extract_number(value) if value else 0
            elif '试验实际负重' in param_name or '实际负载重量' in param_name:
                data['actual_load_weight'] = extract_number(value) if value else 0
            elif '最大设计重量' in param_name:
                data['max_design_weight'] = extract_number(value) if value else 0
            elif '试验运行速度' in param_name or '实际负载速度' in param_name:
                data['actual_load_velocity'] = extract_number(value) if value else 0
            elif '最大设计速度' in param_name:
                data['max_design_velocity'] = extract_number(value) if value else 0
            elif param_name == '机床类型':
                temp_machine_type = str(value).strip() if value else None
                data['load_machine_type'] = temp_machine_type
            elif '评价方案' in param_name:
                temp_evaluation_scheme = str(value).strip() if value else None
                data['evaluation_scheme'] = temp_evaluation_scheme
            elif param_name == '参数1':
                param1_value = extract_number(value) if value else 0
                if temp_machine_type and '铣' in temp_machine_type:
                    data['max_cutting_width'] = param1_value
                elif temp_machine_type and '车' in temp_machine_type:
                    if temp_evaluation_scheme and '径比' in temp_evaluation_scheme:
                        data['max_length_diameter_ratio'] = param1_value
                    else:
                        data['max_cutting_length'] = param1_value
            elif param_name == '参数2':
                param2_value = extract_number(value) if value else 0
                if temp_machine_type and '铣' in temp_machine_type:
                    data['tool_diameter'] = param2_value
                elif temp_machine_type and '车' in temp_machine_type:
                    if temp_evaluation_scheme and '径比' in temp_evaluation_scheme:
                        data['design_length_diameter_ratio'] = param2_value
                    else:
                        data['expected_cutting_length'] = param2_value
        
        if data:
            result = update_machine_parameters_by_type(DamengService, machine_info, data, request)
            if result['success']:
                logger.info(f'成功保存负荷能力数据: {machine_info["name"]}, 数据: {data}')
                success_count += 1
    
    return {'success': True, 'count': success_count}


def _parse_and_save_mtbf(sheet, imported_machines, DamengService, request):
    """
    解析并保存MTBF数据
    
    MTBF = 总运行时间 / 故障次数
    """
    success_count = 0
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    
    for machine_idx, machine_info in enumerate(imported_machines):
        if machine_idx >= len(rows):
            break
            
        row = rows[machine_idx]
        if not row or len(row) < 3:
            continue
        
        data = {}
        total_run_time = extract_number(row[1]) if len(row) > 1 and row[1] else None
        fault_count = int(extract_number(row[2])) if len(row) > 2 and row[2] else None
        
        if total_run_time is not None:
            data['total_run_time'] = total_run_time
        if fault_count is not None:
            data['fault_count'] = fault_count
        
        if data:
            result = update_machine_parameters_by_type(DamengService, machine_info, data, request)
            if result['success']:
                logger.info(f'成功保存MTBF数据: {machine_info["name"]}, 数据: {data}')
                success_count += 1
    
    return {'success': True, 'count': success_count}


def _parse_and_save_repair_time(sheet, imported_machines, DamengService, request):
    """
    解析并保存修复时间数据
    
    用于计算MTTR（平均修复时间）
    """
    success_count = 0
    for machine_idx, machine_info in enumerate(imported_machines):
        repair_times = []
        # 第一个机床的数据在第2列（索引1），第二个机床在第3列（索引2），以此类推
        value_col_idx = 1 + machine_idx
        
        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if not row or len(row) < value_col_idx + 1:
                continue
            
            value = extract_number(row[value_col_idx]) if value_col_idx < len(row) and row[value_col_idx] else 0
            repair_times.append(value)
        
        if repair_times:
            result = update_machine_parameters_by_type(DamengService, machine_info, {'repair_time_values': repair_times}, request)
            if result['success']:
                logger.info(f'成功保存修复时间数据: {machine_info["name"]}, 数据: {repair_times}')
                success_count += 1
    
    return {'success': True, 'count': success_count}


def _parse_and_save_failure_rate(sheet, imported_machines, DamengService, request):
    """
    解析并保存故障率数据
    
    包括致命故障、严重故障、一般故障次数及运行时间
    """
    success_count = 0
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    
    for machine_idx, machine_info in enumerate(imported_machines):
        if machine_idx >= len(rows):
            break
            
        row = rows[machine_idx]
        if not row or len(row) < 5:
            continue
        
        data = {}
        fatal = int(extract_number(row[1])) if len(row) > 1 and row[1] else 0
        serious = int(extract_number(row[2])) if len(row) > 2 and row[2] else 0
        general = int(extract_number(row[3])) if len(row) > 3 and row[3] else 0
        operation_time = extract_number(row[4]) if len(row) > 4 and row[4] else 0
        
        data['fatal_failure'] = fatal
        data['serious_failure'] = serious
        data['general_failure'] = general
        data['operation_time'] = operation_time
        
        if data:
            result = update_machine_parameters_by_type(DamengService, machine_info, data, request)
            if result['success']:
                logger.info(f'成功保存故障率数据: {machine_info["name"]}, 数据: {data}')
                success_count += 1
    
    return {'success': True, 'count': success_count}


def _parse_and_save_consistency(sheet, imported_machines, DamengService, request):
    """
    解析并保存加工一致性数据
    
    包括标准差和均值，用于评估加工稳定性
    """
    success_count = 0
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    
    for machine_idx, machine_info in enumerate(imported_machines):
        if machine_idx >= len(rows):
            break
            
        row = rows[machine_idx]
        if not row or len(row) < 3:
            continue
        
        data = {}
        std_dev = extract_number(row[1]) if len(row) > 1 and row[1] else 0
        mean_val = extract_number(row[2]) if len(row) > 2 and row[2] else 0
        
        data['standard_deviation'] = std_dev
        data['mean_value'] = mean_val
        
        if data:
            result = update_machine_parameters_by_type(DamengService, machine_info, data, request)
            if result['success']:
                logger.info(f'成功保存加工一致性数据: {machine_info["name"]}, 数据: {data}')
                success_count += 1
    
    return {'success': True, 'count': success_count}


def _parse_and_save_precision_retention(sheet, imported_machines, DamengService, request):
    """
    解析并保存加工精度保持性数据
    
    包括原始精度、失效精度、初始精度，用于评估精度衰减
    """
    logger.info(f'开始解析加工精度保持性数据，机床数量: {len(imported_machines)}')
    success_count = 0
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    logger.info(f'读取到 {len(rows)} 行数据')
    
    for machine_idx, machine_info in enumerate(imported_machines):
        if machine_idx >= len(rows):
            logger.warning(f'机床索引 {machine_idx} 超出数据行范围')
            break
            
        row = rows[machine_idx]
        logger.info(f'机床 {machine_info["name"]} (索引{machine_idx}) 行数据: {row[:5] if row and len(row) >= 5 else row}')
        
        if not row or len(row) < 4:
            logger.warning(f'机床 {machine_info["name"]} 行数据不足，跳过')
            continue
        
        data = {}
        original_precision = extract_number(row[1]) if len(row) > 1 and row[1] else 0
        failure_precision = extract_number(row[2]) if len(row) > 2 and row[2] else 0
        initial_precision = extract_number(row[3]) if len(row) > 3 and row[3] else 0
        
        data['original_precision'] = original_precision
        data['failure_precision'] = failure_precision
        data['initial_precision'] = initial_precision
        
        logger.info(f'机床 {machine_info["name"]} 精度保持性数据: original={original_precision}, failure={failure_precision}, initial={initial_precision}')
        
        if data:
            result = update_machine_parameters_by_type(DamengService, machine_info, data, request)
            if result['success']:
                logger.info(f'成功保存加工精度保持性数据: {machine_info["name"]}, 数据: {data}')
                success_count += 1
            else:
                logger.error(f'保存加工精度保持性数据失败: {machine_info["name"]}, 错误: {result.get("message")}')
    
    logger.info(f'加工精度保持性数据导入完成，成功 {success_count} 台')
    return {'success': True, 'count': success_count}


def _parse_and_save_mttf(sheet, imported_machines, DamengService, request):
    """
    解析并保存MTTF数据
    
    MTTF = 首次故障时间，用于评估可靠性
    """
    success_count = 0
    for machine_idx, machine_info in enumerate(imported_machines):
        failure_times = []
        # 第一个机床的数据在第2列（索引1），第二个机床在第3列（索引2），以此类推
        value_col_idx = 1 + machine_idx
        
        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if not row or len(row) < value_col_idx + 1:
                continue
            
            value = extract_number(row[value_col_idx]) if value_col_idx < len(row) and row[value_col_idx] else 0
            failure_times.append(value)
        
        if failure_times:
            result = update_machine_parameters_by_type(DamengService, machine_info, {'failure_time_values': failure_times}, request)
            if result['success']:
                logger.info(f'成功保存MTTF数据: {machine_info["name"]}, 数据: {failure_times}')
                success_count += 1
    
    return {'success': True, 'count': success_count}


def _parse_and_save_intelligence(sheet, imported_machines, DamengService, request):
    """
    解析并保存智能化数据
    
    记录各项智能化特征的满足状态
    """
    logger.info(f'开始解析智能化数据，机床数量: {len(imported_machines)}')
    success_count = 0
    for machine_idx, machine_info in enumerate(imported_machines):
        checkbox_states = []
        # 第一个机床的数据在第4列（索引3），第二个机床在第5列（索引4），以此类推
        satisfied_col_idx = 3 + machine_idx
        
        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if not row or len(row) < satisfied_col_idx + 1:
                logger.info(f'智能化行{row_idx}: 跳过空行或数据不足, row={row}')
                continue
            
            core_feature = str(row[2]).strip() if len(row) > 2 and row[2] else ''
            is_satisfied = str(row[satisfied_col_idx]).strip() if satisfied_col_idx < len(row) and row[satisfied_col_idx] else '0'
            
            logger.info(f'智能化行{row_idx}: C列(核心特征)="{core_feature}", 列{satisfied_col_idx+1}(是否满足)="{is_satisfied}"')
            
            if core_feature:
                satisfied = is_satisfied == '1' or is_satisfied.lower() == 'true' or is_satisfied == '是'
                checkbox_states.append({
                    'feature': core_feature,
                    'satisfied': satisfied
                })
                logger.info(f'智能化行{row_idx}: 添加状态 feature="{core_feature}", satisfied={satisfied}')
        
        if checkbox_states:
            result = update_machine_parameters_by_type(DamengService, machine_info, {'intelligence_checkbox_states': checkbox_states}, request)
            if result['success']:
                logger.info(f'成功保存智能化数据: {machine_info["name"]}, 数据: {checkbox_states}')
                success_count += 1
    
    return {'success': True, 'count': success_count}


def _parse_and_save_resource(sheet, imported_machines, DamengService, request):
    """
    解析并保存资源循环与耗材数据
    
    包括切削液消耗量、润滑油消耗量、轻量化设计项目等
    """
    logger.info(f'开始解析资源循环与耗材数据，机床数量: {len(imported_machines)}')
    success_count = 0
    for machine_idx, machine_info in enumerate(imported_machines):
        data = {}
        lightweight_items = []
        
        rows = list(sheet.iter_rows(min_row=1, values_only=True))
        logger.info(f'读取到 {len(rows)} 行数据')
        
        for row_idx, row in enumerate(rows, 1):
            if not row or len(row) < 2:
                logger.debug(f'行{row_idx}: 跳过空行或数据不足')
                continue
            
            # 前7行从第2列开始（索引1），9至17行从第3列开始（索引2）
            if row_idx <= 7:
                value_col_idx = 1 + machine_idx
            else:
                value_col_idx = 2 + machine_idx
            
            if len(row) < value_col_idx + 1:
                continue
            
            param_name = str(row[0]).strip() if len(row) > 0 and row[0] else ''
            value = row[value_col_idx] if value_col_idx < len(row) else None
            
            logger.info(f'行{row_idx}: row={row[:5] if len(row) >= 5 else row}')
            logger.info(f'行{row_idx}: 参数名="{param_name}", 值={value}')
            
            if '切削液消耗量' in param_name or '切削液消耗' in param_name:
                data['cutting_fluid_q'] = extract_number(value) if value else 0
                logger.info(f'读取切削液消耗量Q: {data["cutting_fluid_q"]}')
            elif '主轴功率 P1' in param_name or ('主轴功率' in param_name and 'P1' in param_name):
                data['cutting_fluid_p'] = extract_number(value) if value else 0
                logger.info(f'读取主轴功率P1: {data["cutting_fluid_p"]}')
            elif '润滑油消耗基准量' in param_name or '基准量 x' in param_name or '润滑油限定' in param_name or '润滑油消耗限定值' in param_name:
                data['lubricant_x'] = extract_number(value) if value else 0
                logger.info(f'读取润滑油消耗基准量x: {data["lubricant_x"]}')
            elif '单位时间润滑油' in param_name or ('润滑油消耗量' in param_name and '基准' not in param_name):
                data['lubricant_l'] = extract_number(value) if value else 0
                logger.info(f'读取润滑油消耗量L: {data["lubricant_l"]}')
            elif '主轴功率 P2' in param_name or ('主轴功率' in param_name and 'P2' in param_name):
                data['lubricant_p'] = extract_number(value) if value else 0
                logger.info(f'读取主轴功率P2: {data["lubricant_p"]}')
            elif row_idx >= 9:
                tech_item = str(row[1]).strip() if len(row) > 1 and row[1] else ''
                is_satisfied = str(row[value_col_idx]).strip() if value_col_idx < len(row) and row[value_col_idx] else '0'
                if tech_item:
                    satisfied = is_satisfied == '1' or is_satisfied.lower() == 'true' or is_satisfied == '是'
                    lightweight_items.append({
                        "feature": tech_item,
                        "satisfied": satisfied
                    })
                    logger.info(f'读取轻量化设计项: {tech_item}, 是否满足: {satisfied}')
        
        if lightweight_items:
            data['lightweight_items'] = lightweight_items
        
        if data:
            logger.info(f'准备保存资源循环与耗材数据: {machine_info["name"]}, 数据: {data}')
            result = update_machine_parameters_by_type(DamengService, machine_info, data, request)
            if result['success']:
                logger.info(f'成功保存资源循环与耗材数据: {machine_info["name"]}, 数据: {data}')
                success_count += 1
            else:
                logger.error(f'保存资源循环与耗材数据失败: {machine_info["name"]}, 错误: {result.get("message")}')
        else:
            logger.warning(f'未找到资源循环与耗材数据: {machine_info["name"]}')
    
    logger.info(f'资源循环与耗材数据导入完成，成功 {success_count} 台')
    return {'success': True, 'count': success_count}


def _parse_and_save_energy(sheet, imported_machines, DamengService, request):
    """
    解析并保存能源-能效数据
    
    包括待机功率、空运转功率、切削功率、材料去除率等
    """
    logger.info(f'开始解析能源-能效数据，机床数量: {len(imported_machines)}')
    success_count = 0
    for machine_idx, machine_info in enumerate(imported_machines):
        data = {}
        # 第一个机床的数据在第2列（索引1），第二个机床在第3列（索引2），以此类推
        value_col_idx = 1 + machine_idx
        
        rows = list(sheet.iter_rows(min_row=1, values_only=True))
        logger.info(f'读取到 {len(rows)} 行数据')
        
        for row_idx, row in enumerate(rows, 1):
            if not row or len(row) < value_col_idx + 1:
                continue
            
            param_name = str(row[0]).strip() if len(row) > 0 and row[0] else ''
            value = row[value_col_idx] if value_col_idx < len(row) else None
            
            logger.info(f'行{row_idx}: row={row[:5] if len(row) >= 5 else row}')
            logger.info(f'行{row_idx}: 参数名="{param_name}", 值={value}')
            
            if '待机功率' in param_name or '待机' in param_name:
                data['est_value'] = extract_number(value) if value else 0
                logger.info(f'读取待机功率Est: {data["est_value"]}')
            elif '空运转功率' in param_name or '空运转' in param_name:
                data['eu_value'] = extract_number(value) if value else 0
                logger.info(f'读取空运转功率Eu: {data["eu_value"]}')
            elif '切削功率' in param_name:
                data['ec_value'] = extract_number(value) if value else 0
                logger.info(f'读取切削功率Ec: {data["ec_value"]}')
            elif '材料去除率' in param_name or '去除率' in param_name:
                data['n_value'] = extract_number(value) if value else 0
                logger.info(f'读取材料去除率n: {data["n_value"]}')
            elif '切削体积' in param_name or '体积' in param_name:
                data['v_value'] = extract_number(value) if value else 0
                logger.info(f'读取切削体积v: {data["v_value"]}')
        
        if data:
            logger.info(f'准备保存能源-能效数据: {machine_info["name"]}, 数据: {data}')
            result = update_machine_parameters_by_type(DamengService, machine_info, data, request)
            if result['success']:
                logger.info(f'成功保存能源-能效数据: {machine_info["name"]}, 数据: {data}')
                success_count += 1
            else:
                logger.error(f'保存能源-能效数据失败: {machine_info["name"]}, 错误: {result.get("message")}')
        else:
            logger.warning(f'未找到能源-能效数据: {machine_info["name"]}')
    
    logger.info(f'能源-能效数据导入完成，成功 {success_count} 台')
    return {'success': True, 'count': success_count}
